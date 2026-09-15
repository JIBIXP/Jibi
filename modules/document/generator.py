"""
Moteur de génération de documents pour JIBI.

Ce module ne gère PAS la sécurité des chemins.
La sécurité et le confinement sont gérés par tools/documents.py.
"""

from docx import Document
from fpdf import FPDF


class DocumentGenerator:
    """Générateur de documents Word et PDF."""

    @staticmethod
    def generate_word(chemin_complet, titre, contenu):
        """
        Génère un document Word à l'emplacement indiqué.

        Le chemin doit déjà avoir été validé et résolu
        par tools.documents.
        """

        document = Document()

        if titre:
            document.add_heading(titre, level=1)

        for paragraphe in contenu.split("\n"):
            if paragraphe.strip():
                document.add_paragraph(paragraphe)

        document.save(chemin_complet)

    @staticmethod
    def generate_pdf(chemin_complet, titre, contenu):
        """
        Génère un document PDF à l'emplacement indiqué.

        Le chemin doit déjà avoir été validé et résolu
        par tools.documents.
        """

        pdf = FPDF()

        pdf.add_page()

        pdf.set_auto_page_break(
            auto=True,
            margin=15
        )

        if titre:
            pdf.set_font(
                "Helvetica",
                "B",
                16
            )

            pdf.multi_cell(
                0,
                10,
                titre
            )

            pdf.ln(4)

        pdf.set_font(
            "Helvetica",
            "",
            12
        )

        for paragraphe in contenu.split("\n"):
            if paragraphe.strip():
                pdf.multi_cell(
                    0,
                    8,
                    paragraphe
                )

                pdf.ln(2)

        pdf.output(chemin_complet)